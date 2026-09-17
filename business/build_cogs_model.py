# -*- coding: utf-8 -*-
"""Build COGS_model_v1.xlsx — unit economics for three business shapes of the cardiac pipeline.

Sheets: README, Inputs, UnitCost_A (산학 수탁), UnitCost_B (V&V40 패키지), UnitCost_C (케이스당 서비스),
        Annual_PnL (5-year scenario), Sensitivity.
Conventions: blue = hardcoded input, black = formula, green = cross-sheet link, yellow fill = key lever.
All money in KRW. Every number on Inputs is an assumption to be edited; sources noted per row.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F, color="000000"); GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True); HDR = Font(name=F, bold=True, color="FFFFFF"); TITLE = Font(name=F, bold=True, size=13)
YELLOW = PatternFill("solid", fgColor="FFFF00"); GREY = PatternFill("solid", fgColor="D9D9D9"); NAVY = PatternFill("solid", fgColor="1F3864")
KRW = '₩#,##0;(₩#,##0);-'; PCT = '0.0%;(0.0%);-'; HRS = '#,##0.0;(#,##0.0);-'; NUM = '#,##0.00;(#,##0.00);-'; INT = '#,##0;(#,##0);-'
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)


def setw(ws, widths):
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w


def header(ws, row, labels):
    for j, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=t); c.font = HDR; c.fill = NAVY; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, bold=False):
    c = ws[ref]; c.value = value; c.font = Font(name=F, color=font.color, bold=bold or font.bold)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c


# =====================================================================================
# Inputs
# =====================================================================================
inp = wb.active; inp.title = "Inputs"
setw(inp, [46, 18, 12, 62])
put(inp, "A1", "가정치 (Inputs) — 파란 글씨는 전부 수정 가능, 노란 칸은 결과에 가장 크게 영향을 주는 레버", bold=True); inp["A1"].font = TITLE
put(inp, "A2", "통화: 원(KRW). 출처 표기: [추정] = 제(Claude) 자릿수 추정, 검증된 시장 데이터 아님. 실제 견적·급여·장비가로 바꿔 쓸 것.")
header(inp, 3, ["항목", "값", "단위", "출처 / 비고"])
R = {}  # key -> cell ref on Inputs


def row(r, key, label, value, unit, note, fmt=KRW, formula=False, lever=False):
    put(inp, f"A{r}", label); c = put(inp, f"B{r}", value, BLACK if formula else BLUE, fmt, YELLOW if lever else None)
    put(inp, f"C{r}", unit); put(inp, f"D{r}", note); R[key] = f"Inputs!$B${r}"
    for col in "ABCD": inp[f"{col}{r}"].border = BOX


r = 4
put(inp, f"A{r}", "1. 인건비", bold=True); inp[f"A{r}"].fill = GREY; r += 1
row(r, "salary", "엔지니어(석사급) 연봉", 55_000_000, "₩/년", "[추정] 국내 의공학 석사 초임~3년차 대역. 실제 채용 조건으로 교체"); r += 1
row(r, "burden", "부대비용 배수(4대보험·퇴직급여 등)", 1.20, "배", "[추정] 통상 1.15~1.25", fmt=NUM); r += 1
row(r, "bill_h", "연간 청구가능(직접투입) 시간", 1_600, "h/년", "[추정] 연 2,000h 근무 중 80%가 프로젝트 직접투입", fmt=INT); r += 1
row(r, "eng_rate", "엔지니어 시간당 원가", f"={R['salary']}*{R['burden']}/{R['bill_h']}", "₩/h", "→ 연봉 × 부대비용 배수 ÷ 청구가능 시간", formula=True); r += 1
row(r, "pi_rate", "PI/시니어 자문 시간당 원가", 80_000, "₩/h", "[추정] 교수·시니어 엔지니어 자문 단가"); r += 1
row(r, "tech_rate", "기술원(랩테크·프린팅 보조) 시간당 원가", 25_000, "₩/h", "[추정]"); r += 1
r += 1; put(inp, f"A{r}", "2. 재료비 (팬텀 1개 기준)", bold=True); inp[f"A{r}"].fill = GREY; r += 1
row(r, "sil_kg", "실리콘 단가", 200_000, "₩/kg", "[추정] Sylgard 184·Ecoflex 급 소매가 대역"); r += 1
row(r, "sil_use", "팬텀당 실리콘 사용량", 1.1, "kg", "phantom_design.json: 실리콘 0.98 L ≈ 1.06 kg + 손실", fmt=NUM); r += 1
row(r, "pva_kg", "PVA 필라멘트 단가", 60_000, "₩/kg", "[추정]"); r += 1
row(r, "pva_use", "코어 PVA 사용량", 0.20, "kg", "core 173.7 mL ≈ 0.2 kg (PVA 밀도 1.2) — 서포트 포함", fmt=NUM); r += 1
row(r, "pla_kg", "PLA 단가", 30_000, "₩/kg", "[추정]"); r += 1
row(r, "pla_use", "주형 상자+절개 모형 PLA 사용량", 0.35, "kg", "mould_box + 다리 + 절개 모형 1세트", fmt=NUM); r += 1
row(r, "fit_cost", "피팅·튜빙·잡자재", 60_000, "₩/팬텀", "[추정] 3/4″·5/8″ 바브 피팅, 튜빙, 이형제, 진공 탈포 소모품"); r += 1
row(r, "mat_phantom", "팬텀 1개 재료비 합계", f"={R['sil_kg']}*{R['sil_use']}+{R['pva_kg']}*{R['pva_use']}+{R['pla_kg']}*{R['pla_use']}+{R['fit_cost']}", "₩", "→ 위 항목의 합", formula=True); r += 1
r += 1; put(inp, f"A{r}", "3. 장비 (감가상각 → 시간/주 단위 원가)", bold=True); inp[f"A{r}"].fill = GREY; r += 1
row(r, "printer_price", "FDM 프린터 가격", 3_000_000, "₩", "[추정] PVA 듀얼 노즐급"); r += 1
row(r, "printer_life", "프린터 수명(가동 시간)", 5_000, "h", "[추정]", fmt=INT); r += 1
row(r, "printer_h", "프린터 시간당 원가", f"={R['printer_price']}/{R['printer_life']}", "₩/h", "→ 가격 ÷ 수명", formula=True); r += 1
row(r, "ws_price", "워크스테이션 가격", 4_000_000, "₩", "RTX 3060 / 32 GB 급 (사용자 보유)"); r += 1
row(r, "ws_years", "워크스테이션 수명", 3, "년", "[추정]", fmt=INT); r += 1
row(r, "ws_hpy", "워크스테이션 연 가동 시간", 2_000, "h/년", "[추정]", fmt=INT); r += 1
row(r, "ws_h", "계산 시간당 원가", f"={R['ws_price']}/({R['ws_years']}*{R['ws_hpy']})", "₩/h", "→ 가격 ÷ (수명 × 연 가동)", formula=True); r += 1
row(r, "bench_price", "펄스 듀플리케이터 시스템(펌프·센서·유량계)", 150_000_000, "₩", "[추정] 상용 심장 유동 벤치 시스템 1~3억 대역. B 형태에만 필요", lever=True); r += 1
row(r, "bench_years", "벤치 수명", 7, "년", "[추정]", fmt=INT); r += 1
row(r, "bench_wpy", "벤치 연 가동 주수", 40, "주/년", "[추정]", fmt=INT); r += 1
row(r, "bench_week", "벤치 주당 원가(감가상각)", f"={R['bench_price']}/({R['bench_years']}*{R['bench_wpy']})", "₩/주", "→ 가격 ÷ (수명 × 연 가동 주수)", formula=True); r += 1
r += 1; put(inp, f"A{r}", "4. 간접비·예비비", bold=True); inp[f"A{r}"].fill = GREY; r += 1
row(r, "oh", "간접비율 (직접인건비 대비)", 0.25, "%", "[추정] 공간·행정·소프트웨어·보험. 대학 산학과제 간접비 15~25%와 유사", fmt=PCT); r += 1
row(r, "cont", "예비비율 (직접원가 대비)", 0.10, "%", "[추정] 재작업·출력 실패·재캐스팅", fmt=PCT); r += 1
r += 1; put(inp, f"A{r}", "5. 가격·규모 레버 — A. 산학 수탁 과제", bold=True); inp[f"A{r}"].fill = GREY; r += 1
row(r, "A_price", "A 과제 단가", 30_000_000, "₩/과제", "[추정] 국내 산학 과제 2천만~5천만 대역", lever=True); r += 1
row(r, "A_cases", "A 과제당 환자 케이스 수", 3, "건", "[추정]", fmt=INT); r += 1
row(r, "A_nre", "A 과제당 비반복 개발(NRE) 시간", 100, "h", "[추정] 기기·질문에 맞춘 파이프라인 개조", fmt=HRS); r += 1
row(r, "A_pm", "A 과제당 미팅·보고·PM 시간", 20, "h", "[추정]", fmt=HRS); r += 1
r += 1; put(inp, f"A{r}", "6. 가격·규모 레버 — B. V&V40 규제 근거 패키지", bold=True); inp[f"A{r}"].fill = GREY; r += 1
row(r, "B_price", "B 프로그램 단가", 200_000_000, "₩/프로그램", "[추정] 해외 CRO 시뮬레이션 신뢰성 패키지·ISO 5840급 벤치 캠페인 1억~5억 대역", lever=True); r += 1
row(r, "B_geoms", "B 프로그램당 환자 형상(팬텀) 수", 5, "건", "[추정] 형상 변이 커버", fmt=INT); r += 1
row(r, "B_pi", "B 시니어/PI 시간", 80, "h", "[추정]", fmt=HRS); r += 1
row(r, "B_bench_w", "B 벤치 점유 주수", 6, "주", "[추정] 셋업 1 + 시험 4 + 재시험 1", fmt=INT); r += 1
row(r, "B_consult", "B 규제 자문 외주(V&V40 검토)", 5_000_000, "₩", "[추정]"); r += 1
row(r, "B_bench_cons", "B 벤치 소모품(혈액 모사액·튜빙·센서 교정)", 1_500_000, "₩", "[추정]"); r += 1
r += 1; put(inp, f"A{r}", "7. 가격·규모 레버 — C. 케이스당 시술 계획 서비스", bold=True); inp[f"A{r}"].fill = GREY; r += 1
row(r, "C_price", "C 케이스 단가", 1_000_000, "₩/건", "[추정] FEops·HeartFlow류 케이스당 과금 대역(50만~150만)", lever=True); r += 1
row(r, "C_h", "C 케이스당 엔지니어 검토 시간", 3, "h", "[추정] 파이프라인 자동화 후 검토·QA만", fmt=HRS); r += 1
row(r, "C_cloud", "C 케이스당 클라우드·스토리지", 10_000, "₩/건", "[추정]"); r += 1
row(r, "C_fixed", "C 연간 고정비(SaMD 유지·QMS·배상책임보험·규제 대응)", 300_000_000, "₩/년", "[추정] 인허가 유지 체계 최소 규모", lever=True); r += 1
row(r, "C_capex", "C 초기 인허가·임상근거 투자", 500_000_000, "₩", "[추정] 식약처 SaMD + 임상 근거. 5년 상각", lever=True); r += 1
row(r, "C_amort", "C 초기 투자 상각 기간", 5, "년", "[추정]", fmt=INT); r += 1
INP_LAST = r
inp.freeze_panes = "A4"

# =====================================================================================
# UnitCost_A — per case, then per project
# =====================================================================================
A = wb.create_sheet("UnitCost_A"); setw(A, [48, 12, 18, 18, 50])
put(A, "A1", "A. 산학 수탁 — 환자 1케이스 단위원가 → 과제 단위 원가·마진", bold=True); A["A1"].font = TITLE
put(A, "A2", "케이스 = CT → 형상 파이프라인 → 팬텀 1개 → 강체벽 CFD → 보고서. 시간(h)은 파란 입력, 단가는 Inputs 참조.")
header(A, 4, ["활동", "시간(h)", "단가(₩/h)", "원가(₩)", "비고"])
acts = [("CT 확보·익명화·품질 확인", 2, "eng", "기관 CT 또는 공개 데이터"),
        ("세그멘테이션 정련(CT 재분할·QA)", 3, "eng", "ct_refine_lv.py 계열; 수동 보정 포함"),
        ("형상 파이프라인 실행·검증(1~17단계)", 4, "eng", "PIPELINE.md; 대부분 자동, 검증 그림 확인"),
        ("팬텀 제작 감독(출력·주형·경화·코어 용해)", 6, "eng", "실제 경과시간은 3~4일, 투입시간만 계상"),
        ("팬텀 제작 보조(기술원)", 6, "tech", "출력 교체·탈포·세척"),
        ("CFD 케이스 설정·실행·후처리", 12, "eng", "OpenFOAM/svMultiPhysics 강체벽"),
        ("보고서 작성", 8, "eng", "형상·팬텀·CFD 결과·불확실성"),
        ("PI 검토", 2, "pi", "")]
r = 5
for name, h, kind, note in acts:
    put(A, f"A{r}", name); put(A, f"B{r}", h, BLUE, HRS)
    rate = {"eng": R["eng_rate"], "tech": R["tech_rate"], "pi": R["pi_rate"]}[kind]
    put(A, f"C{r}", f"={rate}", GREEN, KRW); put(A, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(A, f"E{r}", note)
    r += 1
A_LAB_FIRST, A_LAB_LAST = 5, r - 1
put(A, f"A{r}", "직접인건비 소계", bold=True); put(A, f"B{r}", f"=SUM(B{A_LAB_FIRST}:B{A_LAB_LAST})", BLACK, HRS); put(A, f"D{r}", f"=SUM(D{A_LAB_FIRST}:D{A_LAB_LAST})", BLACK, KRW); A_LAB_SUB = r; r += 2
header(A, r, ["비인건비 항목", "수량", "단가(₩)", "원가(₩)", "비고"]); r += 1
put(A, f"A{r}", "팬텀 재료(실리콘·PVA·PLA·피팅)"); put(A, f"B{r}", 1, BLUE, INT); put(A, f"C{r}", f"={R['mat_phantom']}", GREEN, KRW); put(A, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(A, f"E{r}", "Inputs §2 합계"); r += 1
put(A, f"A{r}", "프린터 가동(코어+상자+절개 모형)"); put(A, f"B{r}", 45, BLUE, HRS); put(A, f"C{r}", f"={R['printer_h']}", GREEN, KRW); put(A, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(A, f"E{r}", "[추정] 출력 시간 45 h"); r += 1
put(A, f"A{r}", "계산(워크스테이션) 가동"); put(A, f"B{r}", 30, BLUE, HRS); put(A, f"C{r}", f"={R['ws_h']}", GREEN, KRW); put(A, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(A, f"E{r}", "[추정] CFD 3주기 + 메시"); r += 1
A_NL_FIRST, A_NL_LAST = A_LAB_SUB + 3, r - 1
put(A, f"A{r}", "비인건비 소계", bold=True); put(A, f"D{r}", f"=SUM(D{A_NL_FIRST}:D{A_NL_LAST})", BLACK, KRW); A_NL_SUB = r; r += 2
put(A, f"A{r}", "직접원가 (인건비+비인건비)", bold=True); put(A, f"D{r}", f"=D{A_LAB_SUB}+D{A_NL_SUB}", BLACK, KRW); A_DIR = r; r += 1
put(A, f"A{r}", "간접비 (직접인건비 × 간접비율)"); put(A, f"D{r}", f"=D{A_LAB_SUB}*{R['oh']}", BLACK, KRW); A_OH = r; r += 1
put(A, f"A{r}", "예비비 (직접원가 × 예비비율)"); put(A, f"D{r}", f"=D{A_DIR}*{R['cont']}", BLACK, KRW); A_CT = r; r += 1
put(A, f"A{r}", "케이스 1건 매출원가 (COGS)", bold=True); put(A, f"D{r}", f"=D{A_DIR}+D{A_OH}+D{A_CT}", BLACK, KRW, YELLOW, bold=True); A_CASE = r; r += 2

put(A, f"A{r}", "과제 단위 (Inputs §5)", bold=True); A[f"A{r}"].fill = GREY; r += 1
header(A, r, ["항목", "수량", "단가(₩)", "원가(₩)", "비고"]); r += 1
put(A, f"A{r}", "케이스 원가 × 케이스 수"); put(A, f"B{r}", f"={R['A_cases']}", GREEN, INT); put(A, f"C{r}", f"=D{A_CASE}", BLACK, KRW); put(A, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); r += 1
put(A, f"A{r}", "비반복 개발(NRE) 인건비"); put(A, f"B{r}", f"={R['A_nre']}", GREEN, HRS); put(A, f"C{r}", f"={R['eng_rate']}*(1+{R['oh']})", BLACK, KRW); put(A, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(A, f"E{r}", "간접비 포함 단가"); r += 1
put(A, f"A{r}", "미팅·보고·PM 인건비"); put(A, f"B{r}", f"={R['A_pm']}", GREEN, HRS); put(A, f"C{r}", f"={R['eng_rate']}*(1+{R['oh']})", BLACK, KRW); put(A, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); r += 1
put(A, f"A{r}", "과제 매출원가 합계", bold=True); put(A, f"D{r}", f"=SUM(D{r-3}:D{r-1})", BLACK, KRW, YELLOW, bold=True); A_PROJ_COGS = r; r += 1
put(A, f"A{r}", "과제 단가(매출)"); put(A, f"D{r}", f"={R['A_price']}", GREEN, KRW); A_PROJ_REV = r; r += 1
put(A, f"A{r}", "매출총이익"); put(A, f"D{r}", f"=D{A_PROJ_REV}-D{A_PROJ_COGS}", BLACK, KRW); r += 1
put(A, f"A{r}", "매출총이익률", bold=True); put(A, f"D{r}", f"=IF(D{A_PROJ_REV}=0,0,(D{A_PROJ_REV}-D{A_PROJ_COGS})/D{A_PROJ_REV})", BLACK, PCT, YELLOW, bold=True); A_GM = r; r += 1
put(A, f"A{r}", "과제당 총 투입 시간(h)"); put(A, f"D{r}", f"={R['A_cases']}*B{A_LAB_SUB}+{R['A_nre']}+{R['A_pm']}", BLACK, HRS); A_HOURS = r; r += 1
put(A, f"A{r}", "엔지니어 1인 연간 최대 과제 수"); put(A, f"D{r}", f"=IF(D{A_HOURS}=0,0,{R['bill_h']}/D{A_HOURS})", BLACK, NUM); put(A, f"E{r}", "→ 청구가능 시간 ÷ 과제당 시간 (용량 상한)"); A_CAP = r; r += 2
put(A, f"A{r}", "해석: 매출총이익률이 높아 보이지만, 이 형태에서 '이익'은 곧 연구실 인건비·장비 재투자다. 용량 상한(연 과제 수)이 매출을 결정한다.")

# =====================================================================================
# UnitCost_B — V&V40 package
# =====================================================================================
B = wb.create_sheet("UnitCost_B"); setw(B, [50, 12, 18, 18, 52])
put(B, "A1", "B. V&V40 규제 근거 패키지 — 기기 프로그램 1건 원가·마진", bold=True); B["A1"].font = TITLE
put(B, "A2", "범위: 환자 형상 N개 + 짝지어진 팬텀 N개 + 벤치 캠페인 + in silico(FSI/EM) 모델·보정·불확실성 정량화 + 신뢰성 보고서. 시간은 파란 입력.")
header(B, 4, ["활동(엔지니어 시간)", "시간(h)", "단가(₩/h)", "원가(₩)", "비고"])
bacts = [("형상 파이프라인 × N (형상당 9 h)", f"=9*{R['B_geoms']}", "형상 수는 Inputs §6"),
         ("팬텀 제작 감독 × N (팬텀당 6 h)", f"=6*{R['B_geoms']}", ""),
         ("벤치 셋업·교정", 40, ""),
         ("벤치 시험 (N 팬텀 × 3 조건 × 4 h)", f"=12*{R['B_geoms']}", "정상·고박출·저박출 등"),
         ("벤치 데이터 분석", 60, ""),
         ("in silico 모델 구축(메시·BC·정합 메시)", 160, "svMultiPhysics FSI 또는 EM"),
         ("모델 보정(벤치 대비)", 140, ""),
         ("불확실성 정량화·민감도(20~30 런 설정·분석)", 160, "V&V40 신뢰성 요소"),
         ("V&V40 신뢰성 문서·보고서", 240, "가장 큰 항목 — 규제 어휘로 작성"),
         ("독립 QA 검토", 60, ""),
         ("PM·고객 미팅", 60, "")]
r = 5
for name, h, note in bacts:
    put(B, f"A{r}", name); put(B, f"B{r}", h, BLUE if not isinstance(h, str) else BLACK, HRS)
    put(B, f"C{r}", f"={R['eng_rate']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(B, f"E{r}", note); r += 1
B_LAB_FIRST, B_LAB_LAST = 5, r - 1
put(B, f"A{r}", "엔지니어 시간 소계", bold=True); put(B, f"B{r}", f"=SUM(B{B_LAB_FIRST}:B{B_LAB_LAST})", BLACK, HRS, YELLOW); put(B, f"D{r}", f"=SUM(D{B_LAB_FIRST}:D{B_LAB_LAST})", BLACK, KRW); B_LAB_SUB = r; r += 1
put(B, f"A{r}", "기술원 시간 (팬텀당 6 h + 벤치 보조 40 h)"); put(B, f"B{r}", f"=6*{R['B_geoms']}+40", BLACK, HRS); put(B, f"C{r}", f"={R['tech_rate']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); B_TECH = r; r += 1
put(B, f"A{r}", "PI/시니어 시간"); put(B, f"B{r}", f"={R['B_pi']}", GREEN, HRS); put(B, f"C{r}", f"={R['pi_rate']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); B_PI = r; r += 1
put(B, f"A{r}", "직접인건비 합계", bold=True); put(B, f"D{r}", f"=D{B_LAB_SUB}+D{B_TECH}+D{B_PI}", BLACK, KRW); B_LAB_TOT = r; r += 2
header(B, r, ["비인건비 항목", "수량", "단가(₩)", "원가(₩)", "비고"]); r += 1
put(B, f"A{r}", "팬텀 재료 × N"); put(B, f"B{r}", f"={R['B_geoms']}", GREEN, INT); put(B, f"C{r}", f"={R['mat_phantom']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); B_NL_FIRST = r; r += 1
put(B, f"A{r}", "프린터 가동 (팬텀당 45 h)"); put(B, f"B{r}", f"=45*{R['B_geoms']}", BLACK, HRS); put(B, f"C{r}", f"={R['printer_h']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); r += 1
put(B, f"A{r}", "벤치 점유 (주)"); put(B, f"B{r}", f"={R['B_bench_w']}", GREEN, INT); put(B, f"C{r}", f"={R['bench_week']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(B, f"E{r}", "감가상각 배분 — 장비가 없으면 외부 벤치 임차비로 교체"); r += 1
put(B, f"A{r}", "벤치 소모품"); put(B, f"B{r}", 1, BLUE, INT); put(B, f"C{r}", f"={R['B_bench_cons']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); r += 1
put(B, f"A{r}", "계산 가동 (FSI/EM 런 + UQ)"); put(B, f"B{r}", 600, BLUE, HRS); put(B, f"C{r}", f"={R['ws_h']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); put(B, f"E{r}", "[추정] 24 h 런 × 25회. 원가는 작고 달력 시간이 병목"); r += 1
put(B, f"A{r}", "규제 자문 외주"); put(B, f"B{r}", 1, BLUE, INT); put(B, f"C{r}", f"={R['B_consult']}", GREEN, KRW); put(B, f"D{r}", f"=B{r}*C{r}", BLACK, KRW); B_NL_LAST = r; r += 1
put(B, f"A{r}", "비인건비 소계", bold=True); put(B, f"D{r}", f"=SUM(D{B_NL_FIRST}:D{B_NL_LAST})", BLACK, KRW); B_NL_SUB = r; r += 2
put(B, f"A{r}", "직접원가", bold=True); put(B, f"D{r}", f"=D{B_LAB_TOT}+D{B_NL_SUB}", BLACK, KRW); B_DIR = r; r += 1
put(B, f"A{r}", "간접비 (직접인건비 × 간접비율)"); put(B, f"D{r}", f"=D{B_LAB_TOT}*{R['oh']}", BLACK, KRW); B_OH = r; r += 1
put(B, f"A{r}", "예비비 (직접원가 × 예비비율)"); put(B, f"D{r}", f"=D{B_DIR}*{R['cont']}", BLACK, KRW); B_CT = r; r += 1
put(B, f"A{r}", "프로그램 매출원가 (COGS)", bold=True); put(B, f"D{r}", f"=D{B_DIR}+D{B_OH}+D{B_CT}", BLACK, KRW, YELLOW, bold=True); B_COGS = r; r += 1
put(B, f"A{r}", "프로그램 단가(매출)"); put(B, f"D{r}", f"={R['B_price']}", GREEN, KRW); B_REV = r; r += 1
put(B, f"A{r}", "매출총이익"); put(B, f"D{r}", f"=D{B_REV}-D{B_COGS}", BLACK, KRW); r += 1
put(B, f"A{r}", "매출총이익률", bold=True); put(B, f"D{r}", f"=IF(D{B_REV}=0,0,(D{B_REV}-D{B_COGS})/D{B_REV})", BLACK, PCT, YELLOW, bold=True); B_GM = r; r += 1
put(B, f"A{r}", "프로그램당 총 인력 시간(h)"); put(B, f"D{r}", f"=B{B_LAB_SUB}+B{B_TECH}+B{B_PI}", BLACK, HRS); B_HOURS = r; r += 1
put(B, f"A{r}", "엔지니어 3인 팀 연간 최대 프로그램 수"); put(B, f"D{r}", f"=IF(B{B_LAB_SUB}=0,0,3*{R['bill_h']}/B{B_LAB_SUB})", BLACK, NUM); put(B, f"E{r}", "→ 3 × 청구가능 시간 ÷ 프로그램당 엔지니어 시간"); r += 2
put(B, f"A{r}", "해석: 원가의 70% 이상이 인건비이고 그중 최대 항목은 문서화다. 파이프라인 자동화가 줄이는 것은 형상·팬텀 시간(전체의 ~10%)뿐 — 마진의 열쇠는 문서 템플릿 재사용이다.")

# =====================================================================================
# UnitCost_C — per-case service
# =====================================================================================
C = wb.create_sheet("UnitCost_C"); setw(C, [46, 18, 18, 18, 18, 18, 18])
put(C, "A1", "C. 케이스당 시술 계획 서비스 — 변동원가·공헌이익·손익분기", bold=True); C["A1"].font = TITLE
put(C, "A2", "전제: SaMD 인허가 완료 후 운영 단계. 고정비·초기투자가 지배적이라 손익분기 케이스 수가 핵심.")
r = 4
put(C, f"A{r}", "케이스당 변동원가", bold=True); C[f"A{r}"].fill = GREY; r += 1
put(C, f"A{r}", "엔지니어 검토 인건비"); put(C, f"B{r}", f"={R['C_h']}*{R['eng_rate']}", BLACK, KRW); r += 1
put(C, f"A{r}", "클라우드·스토리지"); put(C, f"B{r}", f"={R['C_cloud']}", GREEN, KRW); r += 1
put(C, f"A{r}", "계산 (케이스당 4 h)"); put(C, f"B{r}", f"=4*{R['ws_h']}", BLACK, KRW); r += 1
put(C, f"A{r}", "간접비 (인건비 × 간접비율)"); put(C, f"B{r}", f"=B{r-3}*{R['oh']}", BLACK, KRW); r += 1
put(C, f"A{r}", "케이스당 변동원가 합계", bold=True); put(C, f"B{r}", f"=SUM(B{r-4}:B{r-1})", BLACK, KRW, YELLOW, bold=True); C_VAR = r; r += 1
put(C, f"A{r}", "케이스 단가"); put(C, f"B{r}", f"={R['C_price']}", GREEN, KRW); C_P = r; r += 1
put(C, f"A{r}", "케이스당 공헌이익"); put(C, f"B{r}", f"=B{C_P}-B{C_VAR}", BLACK, KRW); C_CM = r; r += 1
put(C, f"A{r}", "공헌이익률"); put(C, f"B{r}", f"=IF(B{C_P}=0,0,B{C_CM}/B{C_P})", BLACK, PCT); r += 2
put(C, f"A{r}", "연간 고정비 (SaMD 유지·QMS·보험)"); put(C, f"B{r}", f"={R['C_fixed']}", GREEN, KRW); C_FIX = r; r += 1
put(C, f"A{r}", "초기 투자 연 상각"); put(C, f"B{r}", f"={R['C_capex']}/{R['C_amort']}", BLACK, KRW); C_AM = r; r += 1
put(C, f"A{r}", "연간 고정비 합계", bold=True); put(C, f"B{r}", f"=B{C_FIX}+B{C_AM}", BLACK, KRW); C_FIXT = r; r += 1
put(C, f"A{r}", "손익분기 케이스 수 (연)", bold=True); put(C, f"B{r}", f"=IF(B{C_CM}<=0,0,B{C_FIXT}/B{C_CM})", BLACK, INT, YELLOW, bold=True); C_BE = r; r += 2
header(C, r, ["연간 케이스 수 →", 200, 400, 600, 800, 1000, 1500]); C_HDR = r
for j in range(2, 8): C.cell(row=r, column=j).font = Font(name=F, bold=True, color="0000FF"); C.cell(row=r, column=j).fill = NAVY; C.cell(row=r, column=j).font = Font(name=F, bold=True, color="FFFFFF")
r += 1
put(C, f"A{r}", "매출")
for j in range(2, 8): col = get_column_letter(j); put(C, f"{col}{r}", f"={col}${C_HDR}*$B${C_P}", BLACK, KRW)
r += 1; put(C, f"A{r}", "변동원가")
for j in range(2, 8): col = get_column_letter(j); put(C, f"{col}{r}", f"={col}${C_HDR}*$B${C_VAR}", BLACK, KRW)
r += 1; put(C, f"A{r}", "고정비 합계")
for j in range(2, 8): col = get_column_letter(j); put(C, f"{col}{r}", f"=$B${C_FIXT}", BLACK, KRW)
r += 1; put(C, f"A{r}", "영업이익", bold=True)
for j in range(2, 8): col = get_column_letter(j); put(C, f"{col}{r}", f"={col}{r-3}-{col}{r-2}-{col}{r-1}", BLACK, KRW, bold=True)
r += 1; put(C, f"A{r}", "필요 엔지니어 수 (검토 시간 기준)")
for j in range(2, 8): col = get_column_letter(j); put(C, f"{col}{r}", f"={col}${C_HDR}*{R['C_h']}/{R['bill_h']}", BLACK, NUM)
r += 2
put(C, f"A{r}", "해석: 공헌이익률은 높지만 손익분기 케이스 수가 국내 대상 시술 건수의 상당 비율이다. 수가 없이 병원이 부담하는 구조에서는 이 케이스 수에 도달하기 어렵다.")

# =====================================================================================
# Annual_PnL — 5-year scenario
# =====================================================================================
P = wb.create_sheet("Annual_PnL"); setw(P, [44, 16, 16, 16, 16, 16, 44])
put(P, "A1", "연간 손익 시나리오 (A → B 전환, C 제외)", bold=True); P["A1"].font = TITLE
put(P, "A2", "파란 행이 시나리오 레버. 인건비는 가동률과 무관하게 전액 비용(급여) 처리 — 미청구 시간이 곧 손실이라는 현실 반영. 원가 시트의 인건비는 여기서 이중 계상하지 않고 '비인건비 직접원가'만 가져온다.")
years = ["1년차", "2년차", "3년차", "4년차", "5년차"]
header(P, 4, ["항목"] + years + ["비고"])
r = 5
put(P, f"A{r}", "시나리오 레버", bold=True); P[f"A{r}"].fill = GREY; r += 1


def prow(label, vals, fmt=INT, blue=True, note=""):
    global r
    put(P, f"A{r}", label)
    for j, v in enumerate(vals):
        col = get_column_letter(2 + j); put(P, f"{col}{r}", v, BLUE if blue else BLACK, fmt)
    put(P, f"G{r}", note); rr = r; r += 1; return rr


R_A = prow("A 과제 수(건/년)", [2, 3, 3, 2, 2], note="[시나리오] 연구실 단계 → B로 무게 이동")
R_B = prow("B 프로그램 수(건/년)", [0, 0, 1, 2, 3], note="[시나리오] 검증 논문·벤치 확보 후 3년차부터")
R_HC = prow("엔지니어 인원(FTE)", [1, 1.5, 3, 3, 4], fmt=NUM, note="[시나리오] 급여는 전액 비용")
R_TECH = prow("기술원 인원(FTE)", [0, 0.5, 1, 1, 1], fmt=NUM)
R_FIX = prow("기타 고정비(임차·행정·SW·보험, ₩/년)", [30_000_000, 40_000_000, 80_000_000, 90_000_000, 100_000_000], fmt=KRW)
R_BENCH = prow("벤치 장비 보유(1=있음)", [0, 1, 1, 1, 1], note="2년차 구매 가정 → 감가상각 시작")
r += 1
put(P, f"A{r}", "손익", bold=True); P[f"A{r}"].fill = GREY; r += 1


def frow(label, fmls, fmt=KRW, note="", bold=False, fill=None):
    global r
    put(P, f"A{r}", label, bold=bold)
    for j, f in enumerate(fmls):
        col = get_column_letter(2 + j); put(P, f"{col}{r}", f, BLACK, fmt, fill, bold=bold)
    put(P, f"G{r}", note); rr = r; r += 1; return rr


cols = [get_column_letter(2 + j) for j in range(5)]
R_REV_A = frow("매출 — A", [f"={c}{R_A}*{R['A_price']}" for c in cols])
R_REV_B = frow("매출 — B", [f"={c}{R_B}*{R['B_price']}" for c in cols])
R_REV = frow("매출 합계", [f"={c}{R_REV_A}+{c}{R_REV_B}" for c in cols], bold=True)
# non-labor direct costs per project: A = (materials+machine)*cases*(1+cont) ; B = non-labor subtotal*(1+cont)
R_DC_A = frow("비인건비 직접원가 — A", [f"={c}{R_A}*UnitCost_A!$D${A_NL_SUB}*{R['A_cases']}*(1+{R['cont']})" for c in cols], note="재료·프린터·계산 (인건비 제외) + 예비비")
R_DC_B = frow("비인건비 직접원가 — B", [f"={c}{R_B}*UnitCost_B!$D${B_NL_SUB}*(1+{R['cont']})" for c in cols], note="벤치 감가상각 배분은 아래 감가상각과 중복되므로 여기서는 제외하지 않음 — 보수적")
R_PAY = frow("인건비 (엔지니어·기술원 급여)", [f"={c}{R_HC}*{R['salary']}*{R['burden']}+{c}{R_TECH}*{R['tech_rate']}*2000" for c in cols], note="기술원은 시급 × 2,000 h")
R_PI = frow("PI/시니어 자문", [f"=({c}{R_A}*2+{c}{R_B}*{R['B_pi']})*{R['pi_rate']}" for c in cols], note="A 과제당 2 h, B 프로그램당 Inputs 시간")
R_FIXC = frow("기타 고정비", [f"={c}{R_FIX}" for c in cols])
R_DEP = frow("감가상각 (벤치·프린터·워크스테이션)", [f"={c}{R_BENCH}*{R['bench_price']}/{R['bench_years']}+{R['printer_price']}/5+{R['ws_price']}/{R['ws_years']}" for c in cols], note="프린터 5년 정액")
R_COST = frow("비용 합계", [f"={c}{R_DC_A}+{c}{R_DC_B}+{c}{R_PAY}+{c}{R_PI}+{c}{R_FIXC}+{c}{R_DEP}" for c in cols], bold=True)
R_OP = frow("영업이익", [f"={c}{R_REV}-{c}{R_COST}" for c in cols], bold=True, fill=YELLOW)
R_OPM = frow("영업이익률", [f"=IF({c}{R_REV}=0,0,{c}{R_OP}/{c}{R_REV})" for c in cols], fmt=PCT)
R_CUM = frow("누적 영업이익", [f"=B{R_OP}"] + [f"={cols[j-1]}{r}+{cols[j]}{R_OP}" for j in range(1, 5)], bold=True)
r += 1
put(P, f"A{r}", "용량 점검", bold=True); P[f"A{r}"].fill = GREY; r += 1
R_NEED = frow("필요 엔지니어 시간(h)", [f"={c}{R_A}*UnitCost_A!$D${A_HOURS}+{c}{R_B}*UnitCost_B!$B${B_LAB_SUB}" for c in cols], fmt=HRS)
R_AVAIL = frow("가용 엔지니어 시간(h)", [f"={c}{R_HC}*{R['bill_h']}" for c in cols], fmt=HRS)
R_UTIL = frow("가동률", [f"=IF({c}{R_AVAIL}=0,0,{c}{R_NEED}/{c}{R_AVAIL})" for c in cols], fmt=PCT, note="100% 초과 = 인원 부족(시나리오 비현실), 60% 미만 = 유휴 인건비")
r += 1
put(P, f"A{r}", "해석: 영업이익은 거의 전적으로 B 프로그램 수와 인원의 가동률로 결정된다. A만으로는 인건비를 넘기 어렵고, B가 연 2건 이상일 때 흑자 구간에 들어간다(가정치 기준).")

# =====================================================================================
# Sensitivity — B margin vs hours & price; A margin vs price & cases
# =====================================================================================
S = wb.create_sheet("Sensitivity"); setw(S, [34, 16, 16, 16, 16, 16])
put(S, "A1", "민감도 — B 프로그램 매출총이익률: 엔지니어 시간 × 단가", bold=True); S["A1"].font = TITLE
put(S, "A2", "행 = 프로그램당 엔지니어 시간(h), 열 = 프로그램 단가(₩). 비인건비·기술원·PI·간접비·예비비는 UnitCost_B 그대로 (시간에 비례하는 항목은 인건비·간접비).")
hours_opts = [800, 1000, 1200, 1500, 1800]; price_opts = [100_000_000, 150_000_000, 200_000_000, 300_000_000, 400_000_000]
header(S, 4, ["시간(h) \\ 단가(₩)"] + [None] * 5)
for j, p in enumerate(price_opts): c = S.cell(row=4, column=2 + j, value=p); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = NAVY; c.number_format = KRW
for i, h in enumerate(hours_opts):
    rr = 5 + i; put(S, f"A{rr}", h, BLUE, HRS)
    for j in range(5):
        col = get_column_letter(2 + j)
        # COGS(h) = [h*eng + tech + pi + nonlabor]*(1+cont) + (h*eng + tech + pi)*oh
        cogs = (f"(($A{rr}*{R['eng_rate']}+UnitCost_B!$D${B_TECH}+UnitCost_B!$D${B_PI}+UnitCost_B!$D${B_NL_SUB})*(1+{R['cont']})"
                f"+($A{rr}*{R['eng_rate']}+UnitCost_B!$D${B_TECH}+UnitCost_B!$D${B_PI})*{R['oh']})")
        put(S, f"{col}{rr}", f"=IF({col}$4=0,0,({col}$4-{cogs})/{col}$4)", BLACK, PCT)
r = 12
put(S, f"A{r}", "민감도 — A 과제 매출총이익률: 과제당 케이스 수 × 단가", bold=True); S[f"A{r}"].font = TITLE; r += 1
cases_opts = [1, 2, 3, 5, 8]; aprice_opts = [15_000_000, 20_000_000, 30_000_000, 40_000_000, 50_000_000]
header(S, r, ["케이스 수 \\ 단가(₩)"] + [None] * 5)
for j, p in enumerate(aprice_opts): c = S.cell(row=r, column=2 + j, value=p); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = NAVY; c.number_format = KRW
hdr = r; r += 1
for i, n in enumerate(cases_opts):
    rr = r + i; put(S, f"A{rr}", n, BLUE, INT)
    for j in range(5):
        col = get_column_letter(2 + j)
        cogs = f"($A{rr}*UnitCost_A!$D${A_CASE}+({R['A_nre']}+{R['A_pm']})*{R['eng_rate']}*(1+{R['oh']}))"
        put(S, f"{col}{rr}", f"=IF({col}${hdr}=0,0,({col}${hdr}-{cogs})/{col}${hdr})", BLACK, PCT)
r = r + len(cases_opts) + 2
put(S, f"A{r}", "읽는 법: 음수 = 원가가 단가를 넘음. B에서 문서화 시간을 줄이는 것(템플릿 재사용)이 단가 인상과 같은 효과를 낸다.")

# =====================================================================================
# README
# =====================================================================================
D = wb.create_sheet("README", 0); setw(D, [110])
lines = [
    ("매출원가(COGS) 모델 v1 — 심장 디지털 트윈 파이프라인의 세 가지 사업 형태", True),
    ("작성 2026-09-16. 모든 숫자는 가정치이며 [추정]으로 표시된 것은 시장 조사가 아닌 자릿수 추정이다. Inputs 시트의 파란 글씨를 실제 값으로 바꾸면 전체가 재계산된다.", False),
    ("", False),
    ("시트 구성", True),
    ("Inputs — 인건비·재료·장비·간접비·가격 레버. 노란 칸이 결과를 가장 크게 움직이는 값.", False),
    ("UnitCost_A — 산학 수탁 과제: 환자 1케이스(형상+팬텀+CFD+보고서) 원가 → 과제 원가·마진·용량.", False),
    ("UnitCost_B — V&V40 규제 근거 패키지: 기기 프로그램 1건의 활동별 시간·비인건비·마진.", False),
    ("UnitCost_C — 케이스당 시술 계획 서비스: 변동원가·공헌이익·손익분기 케이스 수.", False),
    ("Annual_PnL — 5년 시나리오(A→B 전환). 인건비는 가동률과 무관하게 전액 비용 처리(보수적).", False),
    ("Sensitivity — B(시간×단가), A(케이스 수×단가) 매출총이익률 표.", False),
    ("", False),
    ("색 규칙: 파란 글씨 = 입력값, 검정 = 수식, 초록 = 다른 시트 참조, 노란 칸 = 핵심 레버/결과.", False),
    ("", False),
    ("주의", True),
    ("1. 이 모델은 사업 구조를 보기 위한 추정이지 재무 자문이 아니다. 실제 장비 견적·급여·계약 단가로 바꿔야 의미가 있다.", False),
    ("2. Annual_PnL의 인건비는 급여 전액이다. UnitCost 시트의 시간당 원가와 이중 계상하지 않도록 비인건비 직접원가만 가져온다.", False),
    ("3. C 형태는 인허가(SaMD)와 임상 근거가 전제이며, 그 확률과 기간은 모델에 없다 — 손익분기 케이스 수만 본다.", False),
    ("4. MM-WHS 등 연구용 라이선스 데이터는 상업 서비스에 쓸 수 없다. 데이터 확보 비용은 별도.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = D.cell(row=i, column=1, value=t); c.font = Font(name=F, bold=b, size=12 if (b and i == 1) else 10); c.alignment = Alignment(wrap_text=True, vertical="top")

for ws in wb.worksheets:
    for row_ in ws.iter_rows():
        for c in row_:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None)
wb.save(OUT)
print("saved", OUT)
